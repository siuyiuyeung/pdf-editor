"""Real LibreOffice / Ghostscript / pdf2docx runs. Executed inside the worker image (make test-server)."""

import os
import shutil
import subprocess
import time
from pathlib import Path

import pytest

pytestmark = pytest.mark.skipif(
    shutil.which("soffice") is None or shutil.which("gs") is None,
    reason="needs LibreOffice and Ghostscript — run in the worker image: make test-server",
)

from pdftools.errors import JobError  # noqa: E402
from pdftools.handlers import HANDLERS  # noqa: E402
from pdftools.handlers.proc import run  # noqa: E402
from pdftools.runner import run_job  # noqa: E402
from pdftools.storage import JobStore  # noqa: E402


def pages_and_text(pdf: Path) -> tuple[int, str]:
    import pymupdf

    with pymupdf.open(pdf) as doc:
        return doc.page_count, "".join(page.get_text() for page in doc)


def convert(tmp_path: Path, tool: str, src: Path, opts: dict | None = None) -> Path:
    out_dir = tmp_path / "out"
    out_dir.mkdir(exist_ok=True)
    return HANDLERS[tool](src, out_dir, opts or {})


def text_pdf(path: Path, text: str = "Hello Word") -> None:
    import pymupdf

    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), text, fontsize=14)
    doc.save(path)
    doc.close()


def noisy_pdf(path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, 1000, 1000, os.urandom(1000 * 1000 * 3), False)
    page.insert_image(page.rect, pixmap=pix)
    doc.save(path)
    doc.close()


def test_word_to_pdf(tmp_path):
    import docx

    src = tmp_path / "input.docx"
    d = docx.Document()
    d.add_paragraph("Hello from Word")
    d.save(src)
    count, text = pages_and_text(convert(tmp_path, "word-to-pdf", src))
    assert count == 1 and "Hello from Word" in text


def test_legacy_doc_to_pdf(tmp_path):
    import docx

    modern = tmp_path / "legacy.docx"
    d = docx.Document()
    d.add_paragraph("Legacy format")
    d.save(modern)
    subprocess.run(
        ["soffice", f"-env:UserInstallation={(tmp_path / 'profile').as_uri()}", "--headless",
         "--convert-to", "doc", "--outdir", str(tmp_path / "doc"), str(modern)],
        check=True, timeout=120, capture_output=True,
    )
    src = tmp_path / "input.doc"
    shutil.move(tmp_path / "doc" / "legacy.doc", src)
    assert "Legacy format" in pages_and_text(convert(tmp_path, "word-to-pdf", src))[1]


def test_excel_to_pdf(tmp_path):
    import openpyxl

    src = tmp_path / "input.xlsx"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Hello from Excel"
    wb.save(src)
    assert "Hello from Excel" in pages_and_text(convert(tmp_path, "excel-to-pdf", src))[1]


def test_powerpoint_to_pdf(tmp_path):
    from pptx import Presentation

    prs = Presentation()
    for i in (1, 2):
        prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = f"Slide {i}"
    src = tmp_path / "input.pptx"
    prs.save(src)
    count, text = pages_and_text(convert(tmp_path, "powerpoint-to-pdf", src))
    assert count == 2 and "Slide 2" in text


def test_compress_shrinks_an_image_heavy_pdf(tmp_path):
    src = tmp_path / "input.pdf"
    noisy_pdf(src)
    out = convert(tmp_path, "compress-pdf", src, {"level": "high"})
    assert out.stat().st_size < src.stat().st_size / 2
    assert pages_and_text(out)[0] == 1


def test_compress_returns_an_already_small_pdf_unchanged(tmp_path):
    store = JobStore(tmp_path / "jobs")
    job_id = store.new_id()
    store.create(job_id)
    src = store.input_path(job_id, ".pdf")
    text_pdf(src)
    store.write_meta(job_id, {"tool": "compress-pdf", "name": "doc", "size": src.stat().st_size, "opts": {"level": "low"}})
    result = run_job(job_id, "compress-pdf", store=store)
    assert result["ok"] and result["note"] == "already-compressed"
    assert (store.out_dir(job_id) / result["output"]).read_bytes() == src.read_bytes()


def test_pdf_to_word(tmp_path):
    import docx

    src = tmp_path / "input.pdf"
    text_pdf(src, "Hello Word")
    out = convert(tmp_path, "pdf-to-word", src)
    assert out.suffix == ".docx"
    assert "Hello Word" in "\n".join(p.text for p in docx.Document(out).paragraphs)


@pytest.mark.parametrize("tool", ["compress-pdf", "pdf-to-word"])
def test_password_protected_pdfs_report_encrypted(tmp_path, tool):
    import pymupdf

    src = tmp_path / "input.pdf"
    doc = pymupdf.open()
    doc.new_page().insert_text((72, 72), "secret")
    doc.save(src, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
    with pytest.raises(JobError) as err:
        convert(tmp_path, tool, src)
    assert err.value.code == "ENCRYPTED"


def test_garbage_pdf_is_conversion_failed(tmp_path):
    src = tmp_path / "input.pdf"
    src.write_bytes(b"%PDF-1.4\nthis is not really a pdf\n")
    with pytest.raises(JobError) as err:
        convert(tmp_path, "compress-pdf", src)
    assert err.value.code == "CONVERSION_FAILED"


def test_timeout_kills_the_whole_process_group(tmp_path):
    pidfile = tmp_path / "child.pid"
    started = time.monotonic()
    with pytest.raises(JobError) as err:
        run(["sh", "-c", f"sleep 30 & echo $! > {pidfile}; wait"], timeout_s=1)
    assert err.value.code == "TIMEOUT"
    assert time.monotonic() - started < 10
    child = int(pidfile.read_text())
    time.sleep(0.5)  # the container's init (docker run --init) reaps the killed orphan
    with pytest.raises(ProcessLookupError):
        os.kill(child, 0)


def test_a_hanging_pdf_check_times_out(tmp_path, monkeypatch):
    from pdftools.handlers import pdfcheck

    monkeypatch.setattr(pdfcheck, "CHECK_TIMEOUT_S", 1)
    monkeypatch.setattr(pdfcheck, "check_command", lambda src: ["sh", "-c", "sleep 30"])
    src = tmp_path / "input.pdf"
    text_pdf(src)
    started = time.monotonic()
    with pytest.raises(JobError) as err:
        pdfcheck.ensure_openable(src)
    assert err.value.code == "TIMEOUT"
    assert time.monotonic() - started < 10


def test_an_rtf_doc_converts(tmp_path):
    src = tmp_path / "input.doc"
    src.write_bytes(rb"{\rtf1\ansi Hello from RTF\par}")
    assert "Hello from RTF" in pages_and_text(convert(tmp_path, "word-to-pdf", src))[1]


def test_an_html_xls_converts(tmp_path):
    src = tmp_path / "input.xls"
    src.write_bytes(b"<html><body><table><tr><td>Bank export</td></tr></table></body></html>")
    assert "Bank export" in pages_and_text(convert(tmp_path, "excel-to-pdf", src))[1]
