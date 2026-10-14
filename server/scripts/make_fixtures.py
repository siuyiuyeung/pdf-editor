"""Writes the Office sample files used by the server e2e tests: python server/scripts/make_fixtures.py <dir>."""

import sys
from pathlib import Path

import docx
import openpyxl
from pptx import Presentation


def main(out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)

    document = docx.Document()
    document.add_heading("Sample document", 1)
    document.add_paragraph("Hello from Word.")
    document.save(out / "sample.docx")

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["Item", "Qty"])
    sheet.append(["Apples", 3])
    workbook.save(out / "sample.xlsx")

    deck = Presentation()
    for i in (1, 2):
        deck.slides.add_slide(deck.slide_layouts[0]).shapes.title.text = f"Slide {i}"
    deck.save(out / "sample.pptx")


if __name__ == "__main__":
    main(Path(sys.argv[1]))
